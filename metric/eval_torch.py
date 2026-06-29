from PIL import Image
from Metric_torch import *
from natsort import natsorted
from tqdm import tqdm
import os
import numpy as np
import torch
import warnings
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")
device = torch.device("cuda" if torch.cuda.is_available() else "cpu")

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def write_excel(excel_name='metric.xlsx', worksheet_name='VIF', column_index=0, data=None):
    try:
        workbook = load_workbook(excel_name)
    except FileNotFoundError:
        workbook = Workbook()

    worksheet = workbook.create_sheet(title=worksheet_name) if worksheet_name not in workbook.sheetnames else workbook[
        worksheet_name]

    column = get_column_letter(column_index + 1)
    for i, value in enumerate(data):
        cell = worksheet[column + str(i + 1)]
        cell.value = value

    workbook.save(excel_name)


def evaluation_one(ir_name, vi_name, f_name):
    f_img = Image.open(f_name).convert('L')
    ir_img = Image.open(ir_name).convert('L')
    vi_img = Image.open(vi_name).convert('L')

    f_img_tensor = torch.tensor(np.array(f_img)).float().to(device)
    ir_img_tensor = torch.tensor(np.array(ir_img)).float().to(device)
    vi_img_tensor = torch.tensor(np.array(vi_img)).float().to(device)

    f_img_int = np.array(f_img).astype(np.int32)
    f_img_double = np.array(f_img).astype(np.float32)

    ir_img_int = np.array(ir_img).astype(np.int32)
    ir_img_double = np.array(ir_img).astype(np.float32)

    vi_img_int = np.array(vi_img).astype(np.int32)
    vi_img_double = np.array(vi_img).astype(np.float32)

    CE = CE_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    NMI = NMI_function(ir_img_int, vi_img_int, f_img_int, gray_level=256)
    QNCIE = QNCIE_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    TE = TE_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    EI = EI_function(f_img_tensor)
    Qy = Qy_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    Qcb = Qcb_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    EN = EN_function(f_img_tensor)
    MI = MI_function(ir_img_int, vi_img_int, f_img_int, gray_level=256)
    SF = SF_function(f_img_tensor)
    SD = SD_function(f_img_tensor)
    AG = AG_function(f_img_tensor)
    PSNR = PSNR_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    MSE = MSE_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    VIF = VIF_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    CC = CC_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    SCD = SCD_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    Qabf = Qabf_function(ir_img_double, vi_img_double, f_img_double)
    Nabf = Nabf_function(ir_img_tensor, vi_img_tensor, f_img_tensor)
    SSIM = SSIM_function(ir_img_double, vi_img_double, f_img_double)
    MS_SSIM = MS_SSIM_function(ir_img_double, vi_img_double, f_img_double)

    return CE, NMI, QNCIE, TE, EI, Qy, Qcb, EN, MI, SF, AG, SD, CC, SCD, VIF, MSE, PSNR, Qabf, Nabf, SSIM, MS_SSIM


def resolve_fusion_path(fusion_dir, source_filename):
    exact_path = os.path.join(fusion_dir, source_filename)
    if os.path.exists(exact_path):
        return exact_path

    stem, _ = os.path.splitext(source_filename)
    for ext in ('.png', '.jpg', '.jpeg', '.bmp', '.tif', '.tiff'):
        candidate = os.path.join(fusion_dir, stem + ext)
        if os.path.exists(candidate):
            return candidate
    return exact_path


def to_excel_values(values):
    return [
        x.item() if isinstance(x, torch.Tensor) else float(x) if isinstance(x, (int, float)) else x
        for x in values
    ]


def main():
    with_mean = True
    config = {
        'dataroot': os.path.join(PROJECT_ROOT, 'test_img'),
        'results_root': os.path.join(PROJECT_ROOT, 'test_result'),
        'dataset': 'MSRS_v10',
        'fusion_subdir': 'gray',
        'save_dir': os.path.join(PROJECT_ROOT, 'test_result', 'MSRS_v10_xlsx'),
        'max_images': None,
    }

    ir_dir = os.path.join(config['dataroot'], config['dataset'], 'ir')
    vi_dir = os.path.join(config['dataroot'], config['dataset'], 'vi')
    f_dir = os.path.join(config['results_root'], config['dataset'], config['fusion_subdir'])
    os.makedirs(config['save_dir'], exist_ok=True)

    if not os.path.isdir(ir_dir):
        raise FileNotFoundError(f'IR folder not found: {ir_dir}')
    if not os.path.isdir(vi_dir):
        raise FileNotFoundError(f'VI folder not found: {vi_dir}')
    if not os.path.isdir(f_dir):
        raise FileNotFoundError(f'Fusion gray folder not found: {f_dir}')

    filelist = natsorted(os.listdir(ir_dir))
    if config['max_images'] is not None:
        filelist = filelist[:config['max_images']]

    metric_save_name = os.path.join(config['save_dir'], f'metric_{config["dataset"]}_{config["fusion_subdir"]}.xlsx')

    # The current test.py saves objective-metric images under test_result/<dataset>/gray.
    # Use one pseudo-method column to evaluate that folder directly.
    Method_list = ['.']
    start_index = 0

    for i, Method in enumerate(Method_list[start_index:], start=start_index):
        CE_list = []
        NMI_list = []
        QNCIE_list = []
        TE_list = []
        EI_list = []
        Qy_list = []
        Qcb_list = []
        EN_list = []
        MI_list = []
        SF_list = []
        AG_list = []
        SD_list = []
        CC_list = []
        SCD_list = []
        VIF_list = []
        MSE_list = []
        PSNR_list = []
        Qabf_list = []
        Nabf_list = []
        SSIM_list = []
        MS_SSIM_list = []
        filename_list = ['']
        sub_f_dir = os.path.join(f_dir, Method)
        eval_bar = tqdm(filelist)

        for _, item in enumerate(eval_bar):
            ir_name = os.path.join(ir_dir, item)
            vi_name = os.path.join(vi_dir, item)
            f_name = resolve_fusion_path(sub_f_dir, item)

            if os.path.exists(f_name):
                print(ir_name, vi_name, f_name)
                CE, NMI, QNCIE, TE, EI, Qy, Qcb, EN, MI, SF, AG, SD, CC, SCD, VIF, MSE, PSNR, Qabf, Nabf, SSIM, MS_SSIM = evaluation_one(ir_name, vi_name, f_name)
                CE_list.append(CE)
                NMI_list.append(NMI)
                QNCIE_list.append(QNCIE)
                TE_list.append(TE)
                EI_list.append(EI)
                Qy_list.append(Qy)
                Qcb_list.append(Qcb)
                EN_list.append(EN)
                MI_list.append(MI)
                SF_list.append(SF)
                AG_list.append(AG)
                SD_list.append(SD)
                CC_list.append(CC)
                SCD_list.append(SCD)
                VIF_list.append(VIF)
                MSE_list.append(MSE)
                PSNR_list.append(PSNR)
                Qabf_list.append(Qabf)
                Nabf_list.append(Nabf)
                SSIM_list.append(SSIM)
                MS_SSIM_list.append(MS_SSIM)
                filename_list.append(item)
                eval_bar.set_description("{} | {}".format(Method, item))
            else:
                print(f'Skip {item}: fusion image not found at {f_name}')

        if not CE_list:
            raise RuntimeError(f'No fusion images were evaluated in: {sub_f_dir}')

        if with_mean:
            CE_list.append(torch.tensor(CE_list).mean().item())
            NMI_list.append(torch.tensor(NMI_list).mean().item())
            QNCIE_list.append(torch.tensor(QNCIE_list).mean().item())
            TE_list.append(torch.tensor(TE_list).mean().item())
            EI_list.append(torch.tensor(EI_list).mean().item())
            Qy_list.append(torch.tensor(Qy_list).mean().item())
            Qcb_list.append(torch.tensor(Qcb_list).mean().item())
            EN_list.append(torch.tensor(EN_list).mean().item())
            MI_list.append(torch.tensor(MI_list).mean().item())
            SF_list.append(torch.tensor(SF_list).mean().item())
            AG_list.append(torch.tensor(AG_list).mean().item())
            SD_list.append(torch.tensor(SD_list).mean().item())
            CC_list.append(torch.tensor(CC_list).mean().item())
            SCD_list.append(torch.tensor(SCD_list).mean().item())
            VIF_list.append(torch.tensor(VIF_list).mean().item())
            MSE_list.append(torch.tensor(MSE_list).mean().item())
            PSNR_list.append(torch.tensor(PSNR_list).mean().item())
            Qabf_list.append(np.mean(Qabf_list))
            Nabf_list.append(torch.tensor(Nabf_list).mean().item())
            SSIM_list.append(torch.tensor(SSIM_list).mean().item())
            MS_SSIM_list.append(torch.tensor(MS_SSIM_list).mean().item())
            filename_list.append('mean')

        CE_list.insert(0, '{}'.format(Method))
        NMI_list.insert(0, '{}'.format(Method))
        QNCIE_list.insert(0, '{}'.format(Method))
        TE_list.insert(0, '{}'.format(Method))
        EI_list.insert(0, '{}'.format(Method))
        Qy_list.insert(0, '{}'.format(Method))
        Qcb_list.insert(0, '{}'.format(Method))
        EN_list.insert(0, '{}'.format(Method))
        MI_list.insert(0, '{}'.format(Method))
        SF_list.insert(0, '{}'.format(Method))
        AG_list.insert(0, '{}'.format(Method))
        SD_list.insert(0, '{}'.format(Method))
        CC_list.insert(0, '{}'.format(Method))
        SCD_list.insert(0, '{}'.format(Method))
        VIF_list.insert(0, '{}'.format(Method))
        MSE_list.insert(0, '{}'.format(Method))
        PSNR_list.insert(0, '{}'.format(Method))
        Qabf_list.insert(0, '{}'.format(Method))
        Nabf_list.insert(0, '{}'.format(Method))
        SSIM_list.insert(0, '{}'.format(Method))
        MS_SSIM_list.insert(0, '{}'.format(Method))

        if i == start_index:
            for sheet_name in [
                'CE', 'NMI', 'QNCIE', 'TE', 'EI', 'Qy', 'Qcb', 'EN', 'MI', 'SF', 'AG', 'SD',
                'CC', 'SCD', 'VIF', 'MSE', 'PSNR', 'Qabf', 'Nabf', 'SSIM', 'MS_SSIM'
            ]:
                write_excel(metric_save_name, sheet_name, 0, filename_list)

        metric_columns = [
            ('CE', CE_list),
            ('NMI', NMI_list),
            ('QNCIE', QNCIE_list),
            ('TE', TE_list),
            ('EI', EI_list),
            ('Qy', Qy_list),
            ('Qcb', Qcb_list),
            ('EN', EN_list),
            ('MI', MI_list),
            ('SF', SF_list),
            ('AG', AG_list),
            ('SD', SD_list),
            ('CC', CC_list),
            ('SCD', SCD_list),
            ('VIF', VIF_list),
            ('MSE', MSE_list),
            ('PSNR', PSNR_list),
            ('Qabf', Qabf_list),
            ('Nabf', Nabf_list),
            ('SSIM', SSIM_list),
            ('MS_SSIM', MS_SSIM_list),
        ]
        for sheet_name, values in metric_columns:
            write_excel(metric_save_name, sheet_name, i + 1, to_excel_values(values))

    print(f'Metrics saved to: {metric_save_name}')


if __name__ == '__main__':
    main()
